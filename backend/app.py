import os
import re
import spacy
import uvicorn
import docx
import requests
import spacy
from presidio_analyzer import RecognizerRegistry
from presidio_analyzer.nlp_engine import (
    NlpEngine,
    NlpEngineProvider,
)
# import google.generativeai as genai
from dotenv import load_dotenv
from transformers import pipeline
from fastapi.middleware.cors import CORSMiddleware
from presidio_analyzer import AnalyzerEngine
from presidio_anonymizer import AnonymizerEngine

from pydantic import BaseModel
from fastapi import FastAPI, Request, UploadFile, File
from fastapi import FastAPI, Request
from fastapi.responses import JSONResponse

load_dotenv()
app = FastAPI(root_path=os.environ.get("ROOT_PATH"))
# genai.configure(api_key=os.environ.get("GOOGLE_API_KEY"))
# model = genai.GenerativeModel('gemini-pro')
HUGGINGFACE_KEY = os.environ.get("HUGGINGFACE_KEY")
# pipe = pipeline("fill-mask", model="pranavraj1103/ksp-mask-model")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

class TextItem(BaseModel):
    text: str

def create_nlp_engine_with_spacy(
    model_path: str = "en_core_web_lg",
):
    """
    Instantiate an NlpEngine with a spaCy model
    :param model_path: path to model / model name.
    """
    nlp_configuration = {
        "nlp_engine_name": "spacy",
        "models": [{"lang_code": "en", "model_name": model_path}],
        "ner_model_configuration": {
            "model_to_presidio_entity_mapping": {
                "PER": "PERSON",
                "PERSON": "PERSON",
                "NORP": "NRP",
                "FAC": "FACILITY",
                "LOC": "LOCATION",
                "GPE": "LOCATION",
                "LOCATION": "LOCATION",
                "ORG": "ORGANIZATION",
                "ORGANIZATION": "ORGANIZATION",
                "DATE": "DATE_TIME",
                "TIME": "DATE_TIME",
            },
            "low_confidence_score_multiplier": 0.4,
            "low_score_entity_names": ["ORG", "ORGANIZATION"],
        },
    }

    nlp_engine = NlpEngineProvider(nlp_configuration=nlp_configuration).create_engine()

    registry = RecognizerRegistry()
    registry.load_predefined_recognizers(nlp_engine=nlp_engine)

    return nlp_engine, registry

nlp_engine, registry = create_nlp_engine_with_spacy()

analyzer = AnalyzerEngine(nlp_engine=nlp_engine, registry=registry)
anonymizer = AnonymizerEngine()

@app.get("/")
async def read_root():
    return {"message": "Hello World"}


@app.get("/vocab_thresh_masking")
async def vocab_thresh_masking(text, threshold):
    ner_model = spacy.load("en_core_web_sm")
    doc = ner_model(text)
    word_counts = dict()
    for token in doc:
        word_counts[token.text] = word_counts.get(str(token.text), 0) + 1

    threshold = int(threshold)
    frequent_words = [word for word, count in word_counts.items() if count >= threshold]
    masked_text = []
    pii_locations = []  # List to store (start index, end index, type) tuples
    for i, token in enumerate(doc):
        if str(token.text) in frequent_words:
            masked_text.append(str(token.text))
        else:
            masked_text.append("[MASK]")
            # Potentially masked PII, record location and tentative type (UNKNOWN)
            pii_locations.append((token.idx, token.idx + len(token.text), "UNKNOWN"))
    return " ".join(masked_text), pii_locations


@app.get("/entity_tagger_masking")
async def entity_tagger_masking(text):
    ner_model = spacy.load("en_core_web_sm")
    doc = ner_model(text)
    masked_text = []
    pii_locations = []
    for token in doc:
        if token.ent_type_ == "PERSON":
            masked_text.append("[MASK]")
            pii_locations.append((token.idx, token.idx + len(token.text), "PERSON"))
        elif token.ent_type_ == "LOC":
            masked_text.append("[MASK]")
            pii_locations.append((token.idx, token.idx + len(token.text), "LOCATION"))
        elif token.ent_type_ == "ORG":
            masked_text.append("[MASK]")
            pii_locations.append((token.idx, token.idx + len(token.text), "ORGANIZATION"))
        elif token.ent_type_ == "DATE":
            masked_text.append("[MASK]")
            pii_locations.append((token.idx, token.idx + len(token.text), "DATE"))
        else:
            masked_text.append(token.text)
    return " ".join(masked_text), pii_locations


@app.get("/email_and_phone")
async def identify_email_and_phone(text):
    # use regex to identify emails and phone numbers and mask them
    email_pattern = r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}"
    phone_pattern = r"\b\d{3}[-.]?\d{3}[-.]?\d{4}\b"
    
    # find the location of emails and phone numbers
    pii_locations = []
    for match in re.finditer(email_pattern, text):
        pii_locations.append((match.start(), match.end(), "EMAIL"))
    for match in re.finditer(phone_pattern, text):
        pii_locations.append((match.start(), match.end(), "PHONE NUMBER"))

    # mask the emails and phone numbers
    text = re.sub(email_pattern, "[MASK]", text)
    text = re.sub(phone_pattern, "[MASK]", text)
    return text, pii_locations


@app.get("/anonymize_masked_text")
async def anonymize_masked_text(masked_text):
    # prompt = f"The following text contains Personal Information Identifiers marked with [MASK]: \n```\n{masked_text}\n```\n Please anonymize these Personal Identity Identifiers by replacing the '[MASK]' with random placeholders while preserving the context so that the text can be used for analysis."
    # print(prompt)
    # response = model.generate_content(prompt)
    # return response.text
    API_URL = "https://api-inference.huggingface.co/models/pranavraj1103/ksp-mask-model"
    headers = {"Authorization": f"Bearer {HUGGINGFACE_KEY}"}

    def query(payload):
        response = requests.post(API_URL, headers=headers, json=payload)
        return response.json()
        
    output = query({
        "inputs": "The <mask> to the universe is <mask>.",
    })

    return output


@app.post("/parse_doc")
async def parse_doc(file: UploadFile):
    if file.filename.endswith(".txt"):
        return file.file.read()
    doc = docx.Document(file.file)
    full_text = []
    for para in doc.paragraphs:
        full_text.append(para.text)
    return "\n".join(full_text)


@app.post("/presidio_mask")
async def presidio_mask(text: TextItem):
    text = text.text
    results = analyzer.analyze(text=text, language='en')
    # for rec in results:
    #     print(rec.start)
    # print(*[text[res.start : res.end] for res in results])
    # anonymized_text = anonymizer.anonymize(text=text,analyzer_results=results)
    # return anonymized_text, results

    return_list = []
    seen_set = set()
    for rec in results:
        if (rec.score < 0.1) or (rec.start, rec.end) in seen_set:
            continue
        return_list.append({
            "start": rec.start,
            "end": rec.end,
            "entity_type": rec.entity_type,
            "text": text[rec.start:rec.end],
            "score": rec.score,
        })
        seen_set.add((rec.start, rec.end))
    return return_list

